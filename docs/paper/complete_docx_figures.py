#!/usr/bin/env python3
"""Put back figures that the Doc's Word export dropped.

The export of the manuscript Doc leaves out some images (the last three on
2026-10-24). For every "Figure N." caption with no image in the three paragraphs
before it, this inserts the figure's PNG from eval/figures above the caption.

    python3 docs/paper/complete_docx_figures.py EXPORTED.docx [OUT.docx]
"""
import os
import re
import sys

import docx
from docx.shared import Inches

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIG = os.path.join(ROOT, "eval", "figures")
FILES = {1: ["fig_system"], 2: ["fig_scene"], 3: ["fig_orbslam"], 4: ["fig_cad"], 5: ["fig_symmetry"],
         6: ["closures_classes"], 7: ["fig_where"], 8: ["recon_3d", "recon_explain"], 9: ["slam_map_3d"],
         10: ["closures_cost"], 11: ["fig_story"], 12: ["fig_nav"], 13: ["fig_gate_plateau"],
         14: ["fig_endstate"], 15: ["fig_recon_sim2"], 16: ["fig_control"], 17: ["fig_contam"],
         18: ["gates_vs_drift"]}

src = sys.argv[1]
out = sys.argv[2] if len(sys.argv) > 2 else src.replace(".docx", " (all figures).docx")
d = docx.Document(src)
ps = d.paragraphs
width = d.sections[0].page_width - d.sections[0].left_margin - d.sections[0].right_margin
added = []
for i, p in enumerate(ps):
    m = re.match(r"Figure (\d+)\.", p.text.strip())
    if not m:
        continue
    n = int(m.group(1))
    if any(q._p.xpath(".//w:drawing") for q in ps[max(0, i - 3):i]):
        continue
    for f in FILES.get(n, []):
        p.insert_paragraph_before().add_run().add_picture(os.path.join(FIG, f + ".png"), width=width)
    added.append(n)
d.save(out)
print("inserted figures", added, "->", out)
